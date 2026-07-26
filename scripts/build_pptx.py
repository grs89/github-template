#!/usr/bin/env python3
import os
import re
import glob
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import subprocess
import tempfile

def extract_title_from_wiki():
    title = ""
    for fpath in ["01-wiki/Home.md", "README.md", "readme.md"]:
        if os.path.exists(fpath):
            with open(fpath, 'r', encoding='utf-8') as f:
                content = f.read()
                m_title = re.search(r'^#\s+(.+)$', content, re.MULTILINE)
                if m_title:
                    title = re.sub(r'<[^>]+>', '', m_title.group(1).strip())
                    break
    return title or os.environ.get("CI_PROJECT_TITLE", "Proyecto")

PROJECT_TITLE = extract_title_from_wiki()

# ── SISTEMA DE DISEÑO VISUAL (HEREDADO) ──────────────────────────────────────
BG   = RGBColor(0x0A,0x0A,0x0F)
BG2  = RGBColor(0x1A,0x1A,0x28)
PRI  = RGBColor(0x7C,0x3A,0xED)
SEC  = RGBColor(0x06,0xB6,0xD4)
ACC  = RGBColor(0xF5,0x9E,0x0B)
SUC  = RGBColor(0x10,0xB9,0x81)
DAN  = RGBColor(0xEF,0x44,0x44)
TXT  = RGBColor(0xF1,0xF5,0xF9)
MUT  = RGBColor(0x94,0xA3,0xB8)
PSF  = RGBColor(0xA7,0x8B,0xFA)
CSF  = RGBColor(0x67,0xE8,0xF9)
GSF  = RGBColor(0x6E,0xE7,0xB7)
ASF  = RGBColor(0xFC,0xD3,0x4D)
RSF  = RGBColor(0xFC,0xA5,0xA5)
BRD  = RGBColor(0x2D,0x2D,0x44)

W = Inches(13.33)
H = Inches(7.5)

def prs_new():
    p = Presentation()
    p.slide_width = W
    p.slide_height = H
    return p

def sl_new(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def bg(sl, c):
    f = sl.background.fill
    f.solid(); f.fore_color.rgb = c

def rect(sl, l,t,w,h, fc=None, bc=None, bw=Pt(1)):
    s = sl.shapes.add_shape(1,l,t,w,h)
    if fc: s.fill.solid(); s.fill.fore_color.rgb = fc
    else:  s.fill.background()
    if bc: s.line.color.rgb = bc; s.line.width = bw
    else:  s.line.fill.background()
    return s

def tb(sl, txt, l,t,w,h, sz=Pt(12), bold=False, c=None, al=PP_ALIGN.LEFT, fn="Calibri"):
    c = c or TXT
    b = sl.shapes.add_textbox(l,t,w,h)
    tf = b.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = al
    r = p.add_run(); r.text = txt
    r.font.size = sz; r.font.bold = bold
    r.font.color.rgb = c; r.font.name = fn
    return b

def tag(sl, txt, l, t, w=Inches(2.0), h=Inches(0.28), bc=PRI, tc=PSF):
    rect(sl,l,t,w,h,fc=BG2,bc=bc,bw=Pt(0.75))
    tb(sl,txt,l+Inches(.07),t+Inches(.03),w-Inches(.14),h-Inches(.06),sz=Pt(9),bold=True,c=tc,al=PP_ALIGN.CENTER)

def bar_top(sl,c=PRI):   rect(sl,0,0,W,Inches(.055),fc=c)

def bar_bot(sl):
    rect(sl,0,H-Inches(.32),W,Inches(.32),fc=BG2,bc=BRD,bw=Pt(.4))
    tb(sl,f"{PROJECT_TITLE}  |  Wiki-to-PPTX Portal",Inches(.3),H-Inches(.29),Inches(7),Inches(.22),sz=Pt(8),c=MUT)

def snum(sl,n,tot):
    tb(sl,f"{n:02d} / {tot}",W-Inches(1.1),H-Inches(.29),Inches(1.0),Inches(.22),sz=Pt(8),c=MUT,al=PP_ALIGN.RIGHT)

def hdr(sl,tag_txt,tag_c,tag_tc,t1,sub):
    tag(sl,tag_txt,Inches(.4),Inches(.18),Inches(2.3),Inches(.28),bc=tag_c,tc=tag_tc)
    tb(sl,t1,Inches(.4),Inches(.52),Inches(11.5),Inches(.6),sz=Pt(28),bold=True,c=TXT)
    tb(sl,sub,Inches(.4),Inches(1.15),Inches(12.3),Inches(.38),sz=Pt(11),c=MUT)

# ── PARSER DE MARKDOWN PERSONALIZADO ──────────────────────────────────────────
def clean_markdown_inline(text):
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text) # Links [Label](url) -> Label
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)        # Negritas
    text = re.sub(r'\*([^*]+)\*', r'\1', text)            # Cursivas
    text = re.sub(r'`([^`]+)`', r'\1', text)              # Backticks
    return text.strip()

def parse_markdown_to_slides(filepath):
    slides_data = []
    if not os.path.exists(filepath):
        return slides_data
        
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    current_slide = None
    in_code = False
    code_lines = []
    code_lang = ""
    in_table = False
    table_rows = []
    
    file_title = os.path.basename(filepath).replace(".md", "").replace("-", " ")
    
    for line in lines:
        stripped = line.strip()
        
        # Ignorar comentarios de Markdown
        if stripped.startswith('<!--') or stripped.startswith('-->') or stripped.endswith('-->'):
            continue
            
        # Detectar bloques de código
        if stripped.startswith('```'):
            if in_code:
                if current_slide:
                    current_slide['items'].append({
                        'type': 'mermaid' if code_lang == 'mermaid' else 'code',
                        'lang': code_lang,
                        'text': '\n'.join(code_lines)
                    })
                in_code = False
                code_lines = []
                code_lang = ""
            else:
                in_code = True
                code_lang = stripped.replace('```', '').strip().lower()
            continue
            
        if in_code:
            code_lines.append(line.rstrip('\n'))
            continue
            
        # H1: Define el título del documento/sección
        h1_match = re.match(r'^#\s+(.+)$', stripped)
        if h1_match:
            file_title = clean_markdown_inline(h1_match.group(1))
            continue
            
        # H2: Define una nueva diapositiva
        h2_match = re.match(r'^##\s+(.+)$', stripped)
        if h2_match:
            if current_slide and len(current_slide['items']) > 0:
                slides_data.append(current_slide)
                
            current_slide = {
                'title': clean_markdown_inline(h2_match.group(1)),
                'items': [],
                'section': file_title
            }
            in_table = False
            table_rows = []
            continue
            
        if not current_slide:
            continue
            
        # Detectar viñetas
        bullet_match = re.match(r'^[-*+]\s+(.+)$', stripped)
        if bullet_match:
            in_table = False
            current_slide['items'].append({
                'type': 'bullet',
                'text': clean_markdown_inline(bullet_match.group(1))
            })
            continue
            
        # Detectar tablas
        if stripped.startswith('|'):
            if re.match(r'^\|[\s:-|-]+\|$', stripped): # Separador de cabeceras
                continue
            cells = [clean_markdown_inline(c) for c in stripped.split('|')[1:-1]]
            if not in_table:
                in_table = True
                table_rows = [cells]
            else:
                table_rows.append(cells)
                
            current_slide['items'] = [x for x in current_slide['items'] if x['type'] != 'table']
            current_slide['items'].append({
                'type': 'table',
                'rows': table_rows
            })
            continue
            
        # Si no encaja en lo anterior, se trata como un párrafo explicativo
        if stripped:
            in_table = False
            current_slide['items'].append({
                'type': 'paragraph',
                'text': clean_markdown_inline(stripped)
            })
        
    if current_slide and len(current_slide['items']) > 0:
        slides_data.append(current_slide)
        
    return slides_data

# ── RENDERIZADO VISUAL EN PPTX ────────────────────────────────────────────────
def create_section_slide(prs, section_title, slide_idx, total_slides):
    sl = sl_new(prs)
    bg(sl, BG)
    bar_top(sl, PRI)
    bar_bot(sl)
    snum(sl, slide_idx, total_slides)
    
    rect(sl, Inches(1), Inches(2.3), Inches(11.33), Inches(2.9), fc=BG2, bc=PRI, bw=Pt(2.0))
    tb(sl, "SECCIÓN DE LA WIKI", Inches(1.5), Inches(2.7), Inches(10.33), Inches(0.4), sz=Pt(14), bold=True, c=PSF, al=PP_ALIGN.CENTER)
    tb(sl, section_title, Inches(1.5), Inches(3.3), Inches(10.33), Inches(1.5), sz=Pt(36), bold=True, c=TXT, al=PP_ALIGN.CENTER)

def create_content_slide(prs, slide_title, section_tag, items, slide_idx, total_slides):
    sl = sl_new(prs)
    bg(sl, BG)
    bar_top(sl, PRI)
    bar_bot(sl)
    snum(sl, slide_idx, total_slides)
    
    desc = ""
    paragraphs = [x for x in items if x['type'] == 'paragraph']
    if paragraphs:
        desc = paragraphs[0]['text']
        if len(desc) > 160:
            desc = desc[:157] + "..."
        items = [x for x in items if x != paragraphs[0]]
        
    hdr(sl, section_tag.upper(), PRI, PSF, slide_title, desc)
    
    bullets = [x for x in items if x['type'] == 'bullet']
    code_blocks = [x for x in items if x['type'] == 'code']
    mermaid_blocks = [x for x in items if x['type'] == 'mermaid']
    tables = [x for x in items if x['type'] == 'table']
    
    if mermaid_blocks:
        mermaid_txt = mermaid_blocks[0]['text']
        try:
            with tempfile.NamedTemporaryFile(mode='w', suffix='.mmd', delete=False) as f_in:
                f_in.write(mermaid_txt)
                mmd_path = f_in.name
            png_path = mmd_path.replace('.mmd', '.png')
            subprocess.run(["mmdc", "-i", mmd_path, "-o", png_path, "-b", "transparent", "-p", "puppeteer-config.json"], check=True, capture_output=True)
            sl.shapes.add_picture(png_path, Inches(2), Inches(1.8), height=Inches(4.9))
        except Exception as e:
            rect(sl, Inches(.5), Inches(1.8), Inches(12.33), Inches(4.9), fc=BG2, bc=DAN, bw=Pt(1.0))
            tb(sl, f"[Error Renderizando Diagrama Mermaid]\n{str(e)}", Inches(.7), Inches(1.95), Inches(11.93), Inches(4.6), sz=Pt(10), c=RSF)
        finally:
            if 'mmd_path' in locals() and os.path.exists(mmd_path): os.remove(mmd_path)
            if 'png_path' in locals() and os.path.exists(png_path): os.remove(png_path)
            
    elif code_blocks:
        code_txt = code_blocks[0]['text']
        lines = code_txt.split('\n')
        if len(lines) > 18:
            code_txt = '\n'.join(lines[:17]) + "\n... (contenido truncado para PPTX)"
            
        rect(sl, Inches(.5), Inches(1.8), Inches(12.33), Inches(4.9), fc=BG2, bc=BRD, bw=Pt(1.0))
        tb(sl, code_txt, Inches(.7), Inches(1.95), Inches(11.93), Inches(4.6), sz=Pt(9.5), c=CSF, fn="Courier New")
        
    elif tables:
        table_rows = tables[0]['rows']
        max_rows = min(len(table_rows), 5)
        for idx, row in enumerate(table_rows[:max_rows]):
            y = Inches(1.8 + idx * 0.95)
            bg_color = BG2 if idx > 0 else RGBColor(0x23, 0x18, 0x3d)
            border_color = PRI if idx == 0 else BRD
            
            rect(sl, Inches(.5), y, Inches(12.33), Inches(0.85), fc=bg_color, bc=border_color, bw=Pt(1.0))
            
            cell_w = (12.33 - 0.4) / max(len(row), 1)
            for c_idx, cell in enumerate(row):
                bold = (idx == 0 or c_idx == 0)
                cell_c = SEC if (idx == 0 or c_idx == 0) else TXT
                tb(sl, cell, Inches(.6 + c_idx * cell_w), y + Inches(0.22), Inches(cell_w - 0.2), Inches(0.5), sz=Pt(10), bold=bold, c=cell_c)
                
    else:
        all_text_bullets = [b['text'] for b in bullets]
        other_paras = [p['text'] for p in items if p['type'] == 'paragraph']
        all_bullets = all_text_bullets + other_paras
        
        if not all_bullets:
            return
            
        max_bullets = min(len(all_bullets), 10)
        
        if len(all_bullets) > 5:
            mid = (len(all_bullets) + 1) // 2
            col1 = all_bullets[:mid]
            col2 = all_bullets[mid:max_bullets]
            
            col1_txt = "\n\n".join([f"• {b}" for b in col1])
            tb(sl, col1_txt, Inches(.5), Inches(1.9), Inches(5.8), Inches(4.8), sz=Pt(12), c=TXT)
            
            col2_txt = "\n\n".join([f"• {b}" for b in col2])
            tb(sl, col2_txt, Inches(6.8), Inches(1.9), Inches(5.8), Inches(4.8), sz=Pt(12), c=TXT)
        else:
            bullet_txt = "\n\n".join([f"• {b}" for b in all_bullets[:max_bullets]])
            tb(sl, bullet_txt, Inches(.5), Inches(1.9), Inches(12.33), Inches(4.8), sz=Pt(13.5), c=TXT)

# ── TEMPLATE HTML BASE ────────────────────────────────────────────────────────
HTML_TEMPLATE_HEADER = """<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>__PROJECT_TITLE__ — Wiki Presentación</title>
<script src="https://cdn.jsdelivr.net/npm/mermaid@9.4.3/dist/mermaid.min.js"></script>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;600;700;900&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
<style>
:root{--primary:#7C3AED;--secondary:#06B6D4;--accent:#F59E0B;--success:#10B981;--danger:#EF4444;--bg:#0A0A0F;--bg2:#111118;--bg3:#1A1A28;--glass:rgba(255,255,255,0.04);--glass-border:rgba(255,255,255,0.08);--text:#F1F5F9;--text-muted:#94A3B8;--g1:linear-gradient(135deg,#7C3AED 0%,#06B6D4 100%);--themeBorder:#2d2d44;--themeCodeBg:#1A1A28;--themeSecText:#a78bfa}
*{margin:0;padding:0;box-sizing:border-box}
body{font-family:'Inter',sans-serif;background:var(--bg);color:var(--text);overflow:hidden;height:100vh;width:100vw}
.slides{width:100%;height:100%;position:relative}
.slide{position:absolute;inset:0;overflow:hidden;transform:translateX(100%);transition:transform 0.4s cubic-bezier(0.4, 0, 0.2, 1);display:flex;flex-direction:column}
.slide.active{transform:translateX(0);z-index:10}
.nav{position:fixed;bottom:24px;left:50%;transform:translateX(-50%);z-index:100;display:flex;align-items:center;gap:14px;background:rgba(10,10,15,.9);backdrop-filter:blur(20px);border:1px solid var(--glass-border);border-radius:50px;padding:10px 22px}
.nav button{background:none;border:none;color:var(--text-muted);cursor:pointer;font-size:18px;width:34px;height:34px;border-radius:50%;display:flex;align-items:center;justify-content:center;transition:all .2s}
.nav button:hover{background:var(--primary);color:#fff}
#slide-counter{font-size:13px;color:var(--text-muted);font-weight:600;min-width:55px;text-align:center}
.nav-dots{display:flex;gap:5px;flex-wrap:wrap;max-width:240px;justify-content:center}
.dot{width:6px;height:6px;border-radius:50%;background:var(--glass-border);cursor:pointer;transition:all .2s}
.dot.active{background:var(--primary);transform:scale(1.5)}
.tag{display:inline-flex;align-items:center;gap:5px;padding:4px 12px;border-radius:20px;font-size:11px;font-weight:600;letter-spacing:.5px;text-transform:uppercase}
.tag-purple{background:rgba(124,58,237,.2);color:#A78BFA;border:1px solid rgba(124,58,237,.3)}
.card{background:var(--glass);border:1px solid var(--glass-border);border-radius:14px;padding:20px;transition:all .3s}
.card:hover{border-color:rgba(124,58,237,.4);background:rgba(124,58,237,.05)}
.grid-2{display:grid;grid-template-columns:1fr 1fr;gap:24px;width:100%;margin-top:20px}
.si{padding:44px 60px;flex:1;overflow-y:auto;scrollbar-width:none}
.si::-webkit-scrollbar{display:none}
.sh{margin-bottom:28px}
.sh h1{font-size:36px;font-weight:900;line-height:1.1;margin-bottom:8px}
.sh p{color:var(--text-muted);font-size:15px;max-width:850px;line-height:1.6}
.gt{background:var(--g1);-webkit-background-clip:text;-webkit-text-fill-color:transparent;background-clip:text}
.fl{list-style:none;display:flex;flex-direction:column;gap:12px;margin-top:20px}
.fl li{display:flex;align-items:flex-start;gap:11px;font-size:15px;line-height:1.6;color:var(--text-muted)}
.fl li span:first-child{font-size:18px;color:var(--primary);flex-shrink:0;margin-top:1px}
.fl li strong{color:var(--text)}
.snd{position:fixed;top:24px;right:36px;font-size:12px;color:var(--text-muted);font-family:'JetBrains Mono',monospace;font-weight:600;letter-spacing:1px}
.ctag{position:fixed;top:24px;left:36px;display:flex;align-items:center;gap:8px;font-size:12px;color:var(--text-muted);font-weight:600}
.dl{width:8px;height:8px;background:var(--success);border-radius:50%;animation:blink 2s ease-in-out infinite}
@keyframes blink{0%,100%{opacity:1}50%{opacity:.2}}
.kh{position:fixed;top:24px;left:50%;transform:translateX(-50%);font-size:11px;color:var(--text-muted);background:var(--glass);border:1px solid var(--glass-border);border-radius:8px;padding:4px 12px;display:flex;gap:8px;align-items:center}
.kb{background:rgba(255,255,255,.08);border:1px solid rgba(255,255,255,.15);border-radius:4px;padding:1px 6px;font-family:'JetBrains Mono',monospace;font-size:10px}
.meta{display:flex;gap:22px;justify-content:center;flex-wrap:wrap;margin-top:36px}
.mi{display:flex;flex-direction:column;align-items:center;gap:4px;background:var(--bg3);border:1px solid var(--themeBorder);padding:14px 28px;border-radius:12px;min-width:180px}
.mn{font-size:32px;font-weight:900;background:var(--g1);-webkit-background-clip:text;-webkit-text-fill-color:transparent;background-clip:text}
.ml{font-size:11px;color:var(--text-muted);font-weight:600;text-transform:uppercase}
.section-card{background:var(--bg3);border:2px solid var(--primary);border-radius:18px;padding:48px;max-width:700px;width:100%;margin:0 auto;text-align:center;box-shadow:0 10px 30px rgba(124,58,237,0.15)}
.section-card h2{font-size:38px;font-weight:900;color:var(--text);margin-top:14px;line-height:1.2}
.code-container{background:var(--themeCodeBg);border:1px solid var(--themeBorder);border-radius:14px;padding:24px;margin-top:20px;overflow-x:auto;max-height:480px}
.code-container pre{font-family:'JetBrains Mono',monospace;font-size:13px;line-height:1.6;color:var(--text-muted)}
.mermaid{background:var(--themeCodeBg);border:1px solid var(--themeBorder);border-radius:14px;padding:24px;margin-top:20px;display:flex;justify-content:center;width:100%;max-height:480px;overflow:auto}
table{width:100%;border-collapse:collapse;font-size:14px;margin-top:20px;border:1px solid var(--themeBorder);border-radius:12px;overflow:hidden}
th{background:rgba(124,58,237,.15);color:#A78BFA;font-weight:700;text-transform:uppercase;font-size:12px;letter-spacing:.5px;padding:14px 18px;text-align:left;border-bottom:1px solid var(--themeBorder)}
td{padding:12px 18px;border-bottom:1px solid var(--themeBorder);color:var(--text-muted)}
td:first-child{color:var(--text);font-weight:600;width:30%}
tr:last-child td{border-bottom:none}
tr:hover td{background:rgba(124,58,237,.04)}
#s1{background:radial-gradient(ellipse at 30% 40%,rgba(124,58,237,.25) 0%,transparent 60%),radial-gradient(ellipse at 75% 70%,rgba(6,182,212,.15) 0%,transparent 50%),var(--bg);justify-content:center;align-items:center;text-align:center;padding:60px}
#s1 .lr{width:90px;height:90px;border-radius:50%;background:var(--g1);display:flex;align-items:center;justify-content:center;font-size:40px;margin:0 auto 28px;box-shadow:0 0 60px rgba(124,58,237,.5)}
#s1 h1{font-size:52px;font-weight:900;line-height:1;margin-bottom:14px}
#s1 .sub{font-size:18px;color:var(--text-muted);margin-bottom:12px;line-height:1.6}
.sect-div{justify-content:center;align-items:center;padding:60px}
#s_last{background:radial-gradient(ellipse at 25% 50%,rgba(124,58,237,.28) 0%,transparent 60%),radial-gradient(ellipse at 80% 30%,rgba(6,182,212,.18) 0%,transparent 50%),var(--bg);flex:1;display:flex;flex-direction:column;justify-content:center;align-items:center;text-align:center;padding:60px}
#s_last .lr{width:90px;height:90px;border-radius:50%;background:var(--g1);display:flex;align-items:center;justify-content:center;font-size:40px;margin:0 auto 28px;box-shadow:0 0 60px rgba(124,58,237,.5)}
#s_last h1{font-size:52px;font-weight:900;line-height:1;margin-bottom:14px}
#s_last .sub{font-size:18px;color:var(--text-muted);margin-bottom:36px;line-height:1.6}
</style>
</head>
<body>
<div class="ctag"><div class="dl"></div>__PROJECT_TITLE__</div>
<div class="kh"><kbd class="kb">←</kbd><kbd class="kb">→</kbd> Navegar · <kbd class="kb">F</kbd> Pantalla Completa</div>
<div class="snd" id="snd">01 / 18</div>
<div class="slides">
"""

HTML_TEMPLATE_FOOTER = """
</div><!-- /slides -->

<div class="nav">
  <button onclick="prev()">←</button>
  <div class="nav-dots" id="dots"></div>
  <span id="slide-counter">1 / 18</span>
  <button onclick="next()">→</button>
</div>

<script>
try {
  mermaid.initialize({
    startOnLoad: false,
    theme: 'dark',
    themeVariables: {
      primaryColor: '#7C3AED',
      primaryTextColor: '#F1F5F9',
      primaryBorderColor: '#7C3AED',
      lineColor: '#94A3B8',
      secondaryColor: '#1A1A28',
      tertiaryColor: '#0A0A0F',
      background: '#0A0A0F',
      mainBkg: '#1A1A28',
      nodeBorder: '#7C3AED',
      clusterBkg: '#111118',
      titleColor: '#F1F5F9',
      edgeLabelBackground: '#1A1A28',
      fontSize: '12px'
    },
    flowchart: {
      curve: 'basis',
      padding: 14
    }
  });
} catch (e) {
  console.error("Error al inicializar Mermaid:", e);
}

const slides = document.querySelectorAll('.slide');
const dotsC = document.getElementById('dots');
const counter = document.getElementById('slide-counter');
const snd = document.getElementById('snd');
let cur = 0;
const T = slides.length;

if (dotsC) {
  slides.forEach((_, i) => {
    const d = document.createElement('div');
    d.className = 'dot' + (i === 0 ? ' active' : '');
    d.onclick = () => go(i);
    dotsC.appendChild(d);
  });
}

const renderedSlides = new Set();

function go(n) {
  if (slides.length === 0) return;
  
  slides[cur].classList.remove('active');
  if (dotsC && dotsC.children[cur]) {
    dotsC.children[cur].classList.remove('active');
  }
  
  cur = (n + T) % T;
  
  slides[cur].classList.add('active');
  if (dotsC && dotsC.children[cur]) {
    dotsC.children[cur].classList.add('active');
  }
  
  if (counter) counter.textContent = (cur + 1) + ' / ' + T;
  if (snd) snd.textContent = String(cur + 1).padStart(2, '0') + ' / ' + T;
  
  if (!renderedSlides.has(cur)) {
    const mermaidDivs = slides[cur].querySelectorAll('.mermaid');
    if (mermaidDivs.length > 0) {
      setTimeout(() => {
        try {
          mermaid.init(undefined, mermaidDivs);
        } catch (err) {
          console.error("Error renderizando diagrama Mermaid en el slide " + cur + ":", err);
        }
      }, 50);
    }
    renderedSlides.add(cur);
  }
}

go(0);

function next() { go(cur + 1); }
function prev() { go(cur - 1); }

document.addEventListener('keydown', e => {
  if (e.key === 'ArrowRight' || e.key === ' ') next();
  if (e.key === 'ArrowLeft') prev();
  if (e.key === 'f' || e.key === 'F') {
    if (!document.fullscreenElement) document.documentElement.requestFullscreen();
    else document.exitFullscreen();
  }
});
</script>
</body>
</html>
"""

def generate_html_slide(slide_spec, idx, total_slides):
    html = ""
    # Diapositiva Portada
    if idx == 1:
        html += f"""
<!-- SLIDE 1: PORTADA PRINCIPAL -->
<div class="slide active" id="s1">
  <div class="lr">📖</div>
  <h1><span class="gt">{PROJECT_TITLE}</span><br>Documentación Oficial</h1>
  <p class="sub">Documentación Oficial del Proyecto</p>
  <p style="color:var(--text-muted);font-size:14px;">Presentación interactiva autogenerada a partir de los documentos técnicos del proyecto</p>
  
  <div class="meta">
    <div class="mi"><span class="mn">{slide_spec['modules_count']}</span><span class="ml">Módulos Wiki</span></div>
    <div class="mi"><span class="mn">{total_slides - 2}</span><span class="ml">Capítulos de Contenido</span></div>
    <div class="mi"><span class="mn">Auto</span><span class="ml">Generación</span></div>
    <div class="mi"><span class="mn">V2</span><span class="ml">Formato Wiki</span></div>
  </div>
</div>
"""
        return html

    # Diapositiva Cierre
    if idx == total_slides:
        html += f"""
<!-- SLIDE FINAL: CIERRE -->
<div class="slide" id="s_last">
  <div class="lr">📖</div>
  <h1><span class="gt">Fin de la Documentación</span></h1>
  <p class="sub">¡Gracias por leer!</p>
  <p style="color:var(--text-muted);font-size:14px;max-width:550px;line-height:1.6;margin: 0 auto 30px;">Esta presentación se reconstruye automáticamente tras cada commit y cambio en la Wiki, manteniendo el alineamiento de información en tiempo real.</p>
  
  <div class="meta" style="margin-top:0px;">
    <div class="mi"><span class="mn">{slide_spec['modules_count']}</span><span class="ml">Módulos Wiki</span></div>
    <div class="mi"><span class="mn">{total_slides - 2}</span><span class="ml">Capítulos de Contenido</span></div>
    <div class="mi"><span class="mn">Auto</span><span class="ml">Generación</span></div>
    <div class="mi"><span class="mn">V2</span><span class="ml">Formato Wiki</span></div>
  </div>
</div>
"""
        return html

    # Diapositiva de Separación de Sección
    if slide_spec['type'] == 'section':
        html += f"""
<!-- SLIDE {idx}: DIVISOR DE SECCIÓN -->
<div class="slide" id="s{idx}">
  <div class="si sect-div">
    <div class="section-card">
      <div style="font-size:42px;">📖</div>
      <h2 style="font-size:13px;text-transform:uppercase;letter-spacing:1px;color:var(--primary);margin-top:10px;">Sección de la Wiki</h2>
      <h2>{slide_spec['title']}</h2>
    </div>
  </div>
</div>
"""
        return html

    # Diapositiva de Contenido Normal
    title = slide_spec['title']
    section = slide_spec['section']
    items = slide_spec['items']
    
    desc = ""
    paragraphs = [x for x in items if x['type'] == 'paragraph']
    if paragraphs:
        desc = paragraphs[0]['text']
        items = [x for x in items if x != paragraphs[0]]

    html += f"""
<!-- SLIDE {idx}: CONTENIDO -->
<div class="slide" id="s{idx}">
  <div class="si">
    <div class="sh">
      <span class="tag tag-purple">{section}</span>
      <h1 style="margin-top: 6px;">{title}</h1>
      {f'<p>{desc}</p>' if desc else ''}
    </div>
"""
    
    bullets = [x for x in items if x['type'] == 'bullet']
    mermaid_blocks = [x for x in items if x['type'] == 'mermaid']
    code_blocks = [x for x in items if x['type'] == 'code']
    tables = [x for x in items if x['type'] == 'table']
    
    # Caso A: Diagrama Mermaid
    if mermaid_blocks:
        mermaid_txt = mermaid_blocks[0]['text']
        html += f"""
    <div class="mermaid">
{mermaid_txt}
    </div>
"""
    # Caso B: Bloque de Código
    elif code_blocks:
        code_txt = code_blocks[0]['text']
        html += f"""
    <div class="code-container">
      <pre><code>{code_txt}</code></pre>
    </div>
"""

    # Caso B: Tabla
    elif tables:
        table_rows = tables[0]['rows']
        html += "\n    <table>\n"
        for r_idx, row in enumerate(table_rows):
            if r_idx == 0:
                html += "      <thead>\n        <tr>\n"
                for cell in row:
                    html += f"          <th>{cell}</th>\n"
                html += "        </tr>\n      </thead>\n      <tbody>\n"
            else:
                html += "        <tr>\n"
                for cell in row:
                    html += f"          <td>{cell}</td>\n"
                html += "        </tr>\n"
        html += "      </tbody>\n    </table>\n"

    # Caso C: Viñetas
    else:
        all_text_bullets = [b['text'] for b in bullets]
        other_paras = [p['text'] for p in items if p['type'] == 'paragraph']
        all_bullets = all_text_bullets + other_paras
        
        if all_bullets:
            # Si hay más de 5 viñetas, usamos grid de 2 columnas
            if len(all_bullets) > 5:
                mid = (len(all_bullets) + 1) // 2
                col1 = all_bullets[:mid]
                col2 = all_bullets[mid:]
                
                html += '\n    <div class="grid-2">\n      <ul class="fl">\n'
                for b in col1:
                    html += f"        <li><span>•</span><strong>{b}</strong></li>\n"
                html += '      </ul>\n      <ul class="fl">\n'
                for b in col2:
                    html += f"        <li><span>•</span><strong>{b}</strong></li>\n"
                html += '      </ul>\n    </div>\n'
            else:
                html += '\n    <ul class="fl">\n'
                for b in all_bullets:
                    html += f"        <li><span>•</span><strong>{b}</strong></li>\n"
                html += '    </ul>\n'

    html += """
  </div>
</div>
"""
    return html

# ── ORQUESTADOR PRINCIPAL ─────────────────────────────────────────────────────
def main():
    print("🚀 Iniciando compilación de Wiki a formatos PPTX y HTML...")
    prs = prs_new()
    wiki_dir = os.environ.get("WIKI_SOURCE_DIR", "01-wiki")
    
    # Buscar archivos .md dinámicamente
    documents = []
    if os.path.exists(wiki_dir):
        all_mds = glob.glob(os.path.join(wiki_dir, "*.md"))
        documents = [os.path.basename(p) for p in all_mds]
        
        # Buscar Home.md (o home.md) y colocarlo al principio de la lista
        home_file = None
        for doc in list(documents):
            if doc.lower() == "home.md":
                home_file = doc
                documents.remove(doc)
                break
        
        documents = sorted(documents)
        if home_file:
            documents = [home_file] + documents
    else:
        print(f"⚠️ El directorio de Wiki '{wiki_dir}' no existe.")
    
    all_slides_specs = []
    modules_count = 0
    
    # 1. Parsear todos los archivos Markdown de la Wiki
    for doc in documents:
        path = os.path.join(wiki_dir, doc)
        if not os.path.exists(path):
            print(f"⚠️ Archivo no encontrado: {path}. Omitiendo.")
            continue
            
        slides = parse_markdown_to_slides(path)
        if not slides:
            continue
            
        modules_count += 1
        section_name = doc.replace(".md", "").replace("-", " ")
        if slides:
            section_name = slides[0]['section']
            
        # Añadir especificación de diapositiva de divisor de sección
        all_slides_specs.append({
            'type': 'section',
            'title': section_name
        })
        
        # Añadir especificaciones de diapositivas de contenido
        for s in slides:
            all_slides_specs.append({
                'type': 'content',
                'title': s['title'],
                'section': s['section'],
                'items': s['items']
            })
            
    # Limitar a 18 diapositivas de contenido máximo para que la presentación total sea de 20 páginas
    if len(all_slides_specs) > 18:
        print(f"✂️ Truncando presentación de {len(all_slides_specs)} a 18 diapositivas de contenido para respetar el límite de 20 páginas totales.")
        all_slides_specs = all_slides_specs[:18]
        
    # Calcular paginación global
    total_slides = len(all_slides_specs) + 2 # +1 portada +1 cierre
    
    # Portada spec para el generador de HTML
    cover_spec = {'modules_count': modules_count}
    
    # ── GENERACIÓN PPTX ───────────────────────────────────────────────────────
    # Portada PPTX
    sl = sl_new(prs)
    bg(sl, BG)
    bar_top(sl, PRI)
    rect(sl, 0, Inches(.055), W, Inches(.72), fc=BG2)
    rect(sl, Inches(5.9), Inches(.95), Inches(1.45), Inches(1.45), fc=PRI)
    tb(sl, "📖", Inches(5.9), Inches(.98), Inches(1.45), Inches(1.35), sz=Pt(40), al=PP_ALIGN.CENTER)
    tb(sl, PROJECT_TITLE, Inches(1), Inches(2.6), Inches(11.3), Inches(.85), sz=Pt(42), bold=True, c=PSF, al=PP_ALIGN.CENTER)
    tb(sl, "Documentación Oficial del Proyecto", Inches(1), Inches(3.45), Inches(11.3), Inches(.78), sz=Pt(36), bold=True, c=TXT, al=PP_ALIGN.CENTER)
    tb(sl, "Presentación interactiva autogenerada a partir de los documentos técnicos del proyecto", Inches(1.5), Inches(4.3), Inches(10.3), Inches(.55), sz=Pt(13), c=MUT, al=PP_ALIGN.CENTER)
    
    kpi_y = Inches(5.05)
    kpis = [
        (f"{modules_count}", "Módulos Wiki", PSF),
        (f"{len(all_slides_specs)}", "Capítulos de Contenido", CSF),
        ("Auto", "Generación", GSF),
        ("V2", "Formato Wiki", ASF)
    ]
    for idx, (n, l, nc) in enumerate(kpis):
        k_w = Inches(2.4)
        k_h = Inches(1.1)
        x_pos = Inches(1.2 + idx * 2.75)
        rect(sl, x_pos, kpi_y, k_w, k_h, fc=BG2, bc=BRD, bw=Pt(.75))
        tb(sl, n, x_pos, kpi_y + Inches(.12), k_w, Inches(.5), sz=Pt(28), bold=True, c=nc, al=PP_ALIGN.CENTER)
        tb(sl, l, x_pos, kpi_y + Inches(.65), k_w, Inches(.28), sz=Pt(9), bold=True, c=MUT, al=PP_ALIGN.CENTER)
        
    snum(sl, 1, total_slides)
    
    # Diapositivas de contenido y divisores en PPTX
    for s_idx, spec in enumerate(all_slides_specs):
        current_idx = s_idx + 2
        if spec['type'] == 'section':
            create_section_slide(prs, spec['title'], current_idx, total_slides)
        else:
            create_content_slide(prs, spec['title'], spec['section'], spec['items'], current_idx, total_slides)
            
    # Cierre PPTX
    sl = sl_new(prs)
    bg(sl, BG)
    bar_top(sl, PRI)
    rect(sl, 0, Inches(.055), W, Inches(.4), fc=BG2)
    rect(sl, Inches(5.9), Inches(.58), Inches(1.52), Inches(1.52), fc=PRI)
    tb(sl, "📖", Inches(5.9), Inches(.60), Inches(1.52), Inches(1.42), sz=Pt(42), al=PP_ALIGN.CENTER)
    tb(sl, "Fin de la Documentación", Inches(1.0), Inches(2.55), Inches(11.3), Inches(.78), sz=Pt(42), bold=True, c=PSF, al=PP_ALIGN.CENTER)
    tb(sl, "¡Gracias por leer!", Inches(1.0), Inches(3.35), Inches(11.3), Inches(.78), sz=Pt(36), bold=True, c=TXT, al=PP_ALIGN.CENTER)
    tb(sl, "Esta presentación se reconstruye automáticamente tras cada commit y cambio en la Wiki.", Inches(1.5), Inches(4.3), Inches(10.3), Inches(.68), sz=Pt(13), c=MUT, al=PP_ALIGN.CENTER)
    
    for idx, (n, l, nc) in enumerate(kpis):
        x_pos = Inches(1.2 + idx * 2.75)
        rect(sl, x_pos, kpi_y, k_w, k_h, fc=BG2, bc=BRD, bw=Pt(.75))
        tb(sl, n, x_pos, kpi_y + Inches(.12), k_w, Inches(.5), sz=Pt(28), bold=True, c=nc, al=PP_ALIGN.CENTER)
        tb(sl, l, x_pos, kpi_y + Inches(.65), k_w, Inches(.28), sz=Pt(9), bold=True, c=MUT, al=PP_ALIGN.CENTER)
        
    snum(sl, total_slides, total_slides)
    
    # Guardar PPTX
    project_name = os.environ.get("CI_PROJECT_NAME", "Proyecto")
    out_pptx = f"public/Wiki_{project_name}.pptx"
    prs.save(out_pptx)
    print(f"   Presentación PPTX guardada en: {out_pptx}")
    
    # ── GENERACIÓN HTML INTERACTIVA ───────────────────────────────────────────
    html_content = HTML_TEMPLATE_HEADER.replace("__PROJECT_TITLE__", PROJECT_TITLE)
    
    # Diapositiva 1: Portada
    html_content += generate_html_slide(cover_spec, 1, total_slides)
    
    # Diapositivas 2 a N-1
    for s_idx, spec in enumerate(all_slides_specs):
        current_idx = s_idx + 2
        html_content += generate_html_slide(spec, current_idx, total_slides)
        
    # Diapositiva N: Cierre
    html_content += generate_html_slide(cover_spec, total_slides, total_slides)
    
    html_content += HTML_TEMPLATE_FOOTER
    
    # Guardar HTML
    out_html = f"public/Wiki_{project_name}.html"
    with open(out_html, 'w', encoding='utf-8') as f:
        f.write(html_content)
        
    print(f"   Presentación HTML interactiva guardada en: {out_html}")
    print(f"✅ ¡Éxito! Generación dual completada. Total diapositivas: {total_slides}")

if __name__ == "__main__":
    main()
